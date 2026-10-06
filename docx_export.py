"""Turn the audit report's Markdown into a formatted Word document.

Handles only what the report uses: #/##/### headings, paragraphs, bullets, '>' quotes,
pipe tables and **bold**. A line ending in two spaces continues the same paragraph.
"""
import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_inline(par, text):
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        par.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def md_to_docx(md: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = [l.rstrip("\r") for l in md.split("\n")]
    cur = None          # paragraph we may still append hard-broken lines to
    hard_break = False  # previous line ended with two spaces
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        i += 1
        if not stripped:
            cur, hard_break = None, False
            continue

        if stripped.startswith("|"):  # table: gather consecutive pipe rows
            rows = [stripped]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows
                     if not re.fullmatch(r"[|\-\s:]+", r)]
            if cells:
                t = doc.add_table(rows=len(cells), cols=len(cells[0]))
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(cells):
                    for ci, val in enumerate(row[:len(cells[0])]):
                        t.cell(ri, ci).text = val
                doc.add_paragraph()
            cur = None
            continue

        m = re.match(r"(#{1,3}) (.*)", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=0 if level == 1 else level - 1)
            cur, hard_break = None, False
            continue

        if stripped.startswith(">"):
            text = stripped.lstrip("> ").rstrip()
            if cur is not None and cur.style.name == "Quote":
                cur.add_run().add_break(WD_BREAK.LINE)
            else:
                cur = doc.add_paragraph(style="Quote")
            add_inline(cur, text)
            hard_break = False
            continue

        if stripped.startswith("- "):
            cur = doc.add_paragraph(style="List Bullet")
            add_inline(cur, stripped[2:])
            hard_break = line.endswith("  ")
            continue

        # Indented line under a bullet, or the line after a hard break: same paragraph, new line.
        if cur is not None and (line.startswith("  ") or hard_break):
            cur.add_run().add_break(WD_BREAK.LINE)
            add_inline(cur, stripped)
            hard_break = line.endswith("  ")
            continue

        cur = doc.add_paragraph()
        add_inline(cur, stripped)
        hard_break = line.endswith("  ")

    # Footer note so a printed copy still says what it is.
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    run = fp.add_run("Draft prepared with AI assistance. All statuses require review by a qualified person.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
